"""Non-scientific structural/integrity audit of data/raw/tb_mes.xlsx.

Performs ONLY dataset integrity inspection: sheet/row/column shape, parsed date
range, duplicate/missing months, missing values, numeric/non-numeric problems,
and observation counts within the canonical calibration/validation intervals.

Does NOT compute any epidemiological model parameter. Does NOT modify the raw
source file.
"""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[1]
RAW_PATH = REPO_ROOT / "data" / "raw" / "tb_mes.xlsx"
OUTPUT_PATH = REPO_ROOT / "outputs" / "audits" / "dataset_integrity.json"

CALIBRATION_START = date(2001, 1, 1)
CALIBRATION_END = date(2020, 12, 1)
VALIDATION_START = date(2021, 1, 1)
VALIDATION_END = date(2022, 12, 1)

DATE_COL = "data"
YEAR_COL = "ano"
MONTH_COL = "monthes"
NUMERIC_COLS = ["casos", "populacao", "TI"]


def month_range(start: date, end: date) -> list[date]:
    months = []
    y, m = start.year, start.month
    while (y, m) <= (end.year, end.month):
        months.append(date(y, m, 1))
        m += 1
        if m == 13:
            m = 1
            y += 1
    return months


def main() -> None:
    report: dict = {"source_file": str(RAW_PATH.relative_to(REPO_ROOT))}

    wb = openpyxl.load_workbook(RAW_PATH, data_only=True)
    report["sheet_names"] = wb.sheetnames

    ws = wb[wb.sheetnames[0]]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    header_trimmed = [h for h in header if h is not None]
    data_rows = list(ws.iter_rows(min_row=2, values_only=True))

    report["active_sheet"] = ws.title
    report["max_row"] = ws.max_row
    report["max_column"] = ws.max_column
    report["column_names"] = header_trimmed
    report["data_row_count"] = len(data_rows)

    col_index = {name: header.index(name) for name in header if name is not None}

    dates: list[date | None] = []
    missing_value_cells: list[dict] = []
    non_numeric_problems: list[dict] = []

    for i, row in enumerate(data_rows, start=2):
        raw_date = row[col_index[DATE_COL]] if DATE_COL in col_index else None
        d = raw_date.date() if hasattr(raw_date, "date") else raw_date
        dates.append(d)

        for col_name in [DATE_COL, YEAR_COL, MONTH_COL] + NUMERIC_COLS:
            if col_name not in col_index:
                continue
            value = row[col_index[col_name]]
            if value is None:
                missing_value_cells.append({"row": i, "column": col_name})

        for col_name in NUMERIC_COLS:
            if col_name not in col_index:
                continue
            value = row[col_index[col_name]]
            if value is not None and not isinstance(value, (int, float)):
                non_numeric_problems.append(
                    {"row": i, "column": col_name, "value": repr(value)}
                )

    valid_dates = [d for d in dates if d is not None]
    report["parsed_date_min"] = valid_dates[0].isoformat() if valid_dates else None
    report["parsed_date_max"] = valid_dates[-1].isoformat() if valid_dates else None
    report["null_dates_count"] = dates.count(None)

    seen: dict[date, int] = {}
    for d in valid_dates:
        seen[d] = seen.get(d, 0) + 1
    duplicate_dates = sorted(d.isoformat() for d, n in seen.items() if n > 1)
    report["duplicate_dates"] = duplicate_dates
    report["duplicate_date_count"] = len(duplicate_dates)

    if valid_dates:
        expected_months = month_range(min(valid_dates), max(valid_dates))
        present = set(valid_dates)
        missing_months = sorted(
            d.isoformat() for d in expected_months if d not in present
        )
    else:
        missing_months = []
    report["missing_months"] = missing_months
    report["missing_month_count"] = len(missing_months)

    report["missing_value_cells"] = missing_value_cells
    report["missing_value_count"] = len(missing_value_cells)
    report["non_numeric_problems"] = non_numeric_problems
    report["non_numeric_problem_count"] = len(non_numeric_problems)

    calibration_months = month_range(CALIBRATION_START, CALIBRATION_END)
    validation_months = month_range(VALIDATION_START, VALIDATION_END)
    present = set(valid_dates)
    calibration_count = sum(1 for d in calibration_months if d in present)
    validation_count = sum(1 for d in validation_months if d in present)

    report["canonical_total_expected"] = 264
    report["canonical_calibration_expected"] = 240
    report["canonical_validation_expected"] = 24
    report["total_observations_found"] = len(valid_dates)
    report["calibration_interval"] = {
        "start": CALIBRATION_START.isoformat(),
        "end": CALIBRATION_END.isoformat(),
        "observations_found": calibration_count,
    }
    report["validation_interval"] = {
        "start": VALIDATION_START.isoformat(),
        "end": VALIDATION_END.isoformat(),
        "observations_found": validation_count,
    }

    report["matches_canonical_expectation"] = (
        report["total_observations_found"] == 264
        and calibration_count == 240
        and validation_count == 24
    )

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(json.dumps(report, indent=2, ensure_ascii=False))

    print(f"Wrote {OUTPUT_PATH.relative_to(REPO_ROOT)}")
    print(json.dumps(report, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
