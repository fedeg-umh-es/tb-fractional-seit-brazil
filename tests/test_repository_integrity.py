"""Repository-level integrity tests.

These tests only check provenance/manifest bookkeeping and non-scientific
dataset structure (observation counts and interval partitioning). They do not
compute or assert any epidemiological model result.
"""

import csv
import hashlib
from datetime import date
from pathlib import Path

import openpyxl
import pytest

REPO_ROOT = Path(__file__).resolve().parents[1]
RAW_DATASET = REPO_ROOT / "data" / "raw" / "tb_mes.xlsx"
RAW_MANUSCRIPT = REPO_ROOT / "manuscript" / "source" / "BIOMATEMATICA_UNICAMP.docx"
MANIFEST = REPO_ROOT / "outputs" / "audits" / "source_manifest.csv"

CALIBRATION_START = date(2001, 1, 1)
CALIBRATION_END = date(2020, 12, 1)
VALIDATION_START = date(2021, 1, 1)
VALIDATION_END = date(2022, 12, 1)


def sha256_of(path: Path) -> str:
    h = hashlib.sha256()
    h.update(path.read_bytes())
    return h.hexdigest()


def load_manifest_rows() -> list[dict]:
    with MANIFEST.open(newline="") as f:
        return list(csv.DictReader(f))


def load_dataset_dates() -> list[date]:
    wb = openpyxl.load_workbook(RAW_DATASET, data_only=True)
    ws = wb[wb.sheetnames[0]]
    dates = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        raw_date = row[0]
        if raw_date is None:
            continue
        dates.append(raw_date.date() if hasattr(raw_date, "date") else raw_date)
    return dates


def test_raw_source_files_exist():
    assert RAW_DATASET.is_file(), "canonical dataset missing from data/raw/"
    assert RAW_MANUSCRIPT.is_file(), "source manuscript missing from manuscript/source/"


def test_manifest_fields_populated():
    rows = load_manifest_rows()
    assert len(rows) == 2
    required_fields = [
        "artifact",
        "original_filename",
        "repository_path",
        "size_bytes",
        "sha256",
        "role",
        "status",
    ]
    for row in rows:
        for field in required_fields:
            assert row[field], f"manifest field {field!r} empty for {row['artifact']}"
        assert row["status"] == "PRESERVED_RAW_SOURCE"


def test_manifest_hashes_match_files():
    rows = {row["artifact"]: row for row in load_manifest_rows()}
    assert rows["tb_mes.xlsx"]["sha256"] == sha256_of(RAW_DATASET)
    assert (
        rows["BIOMATEMATICA_UNICAMP.docx"]["sha256"] == sha256_of(RAW_MANUSCRIPT)
    )


def test_dataset_has_264_monthly_observations():
    dates = load_dataset_dates()
    assert len(dates) == 264


def test_calibration_interval_has_240_observations():
    dates = load_dataset_dates()
    calibration = [d for d in dates if CALIBRATION_START <= d <= CALIBRATION_END]
    assert len(calibration) == 240


def test_validation_interval_has_24_observations():
    dates = load_dataset_dates()
    validation = [d for d in dates if VALIDATION_START <= d <= VALIDATION_END]
    assert len(validation) == 24


def test_calibration_and_validation_do_not_overlap():
    dates = load_dataset_dates()
    calibration = {d for d in dates if CALIBRATION_START <= d <= CALIBRATION_END}
    validation = {d for d in dates if VALIDATION_START <= d <= VALIDATION_END}
    assert calibration.isdisjoint(validation)
